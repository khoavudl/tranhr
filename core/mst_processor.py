"""
Module xử lý Cấp Mã số thuế cá nhân (MST)
Logic được chuyển đổi nguyên vẹn từ MST.ipynb (Phiên bản 1.2)
"""

import os
import re
from typing import Callable, List, Optional, Dict, Any
try:
    import pymupdf as fitz
except ImportError:
    import fitz
from openpyxl import load_workbook
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter


def read_pdf(path: str) -> str:
    """Đọc toàn bộ text trong PDF với sort=True."""
    text = ""
    doc = fitz.open(path)
    for page in doc:
        text += page.get_text("text", sort=True)
    return text


def parse_14(text: str, pdf_name: str) -> Dict[str, Dict[str, str]]:
    """Trích xuất CCCD được cấp MST mới từ Mẫu 14."""
    results = {}
    cccd_list = re.findall(r'\b\d{12}\b', text)

    for cccd in cccd_list:
        results[cccd] = {
            "mst": cccd,
            "status": "CAP_MOI",
            "note": "Được cấp MST",
            "source": pdf_name
        }
    return results


def parse_40(text: str, pdf_name: str) -> Dict[str, Dict[str, str]]:
    """Trích xuất CCCD và MST cũ từ Mẫu 40/TB-ĐKT."""
    results = {}
    text_clean = text.replace("\n", " ")

    # Regex: Quét qua các số lẻ tẻ giữa chữ MST và mã 10 số nhưng không nhảy qua CCCD 12 số khác
    pattern = r'\b(\d{12})\b(?:(?!\b\d{12}\b).)*?MST(?:(?!\b\d{12}\b).)*?(\b\d{10}\b)'

    matches = re.findall(pattern, text_clean, flags=re.DOTALL)

    for cccd, mst in matches:
        results[cccd] = {
            "mst": mst,
            "status": "TRUNG_MST",
            "note": f"MST đã tồn tại: {mst}",
            "source": pdf_name
        }
    return results


def process_mst(
    excel_file: str,
    pdf_files: List[str],
    output_file: Optional[str] = None,
    log_callback: Optional[Callable[[str], None]] = None,
    progress_callback: Optional[Callable[[int, int], None]] = None
) -> Dict[str, Any]:
    """
    Thực thi quy trình đối soát & cập nhật file MST.
    
    Args:
        excel_file: Đường dẫn file Excel nguồn (Tong hop MST.xlsx)
        pdf_files: Danh sách đường dẫn file PDF thông báo từ Cục Thuế
        output_file: Đường dẫn file xuất kết quả (nếu None sẽ tự sinh)
        log_callback: Hàm nhận log message hiển thị lên UI
        progress_callback: Hàm nhận tiến độ (current, total)
        
    Returns:
        Dict chứa thống kê kết quả
    """
    def log(msg: str):
        if log_callback:
            log_callback(msg)
        else:
            print(msg)

    if not os.path.exists(excel_file):
        raise FileNotFoundError(f"Không tìm thấy file Excel: {excel_file}")
    if not pdf_files:
        raise ValueError("Chưa chọn file PDF nào để đối soát.")

    log(f"Bắt đầu xử lý MST...")
    log(f"- File Excel: {os.path.basename(excel_file)}")
    log(f"- Số lượng PDF: {len(pdf_files)} file")

    all_results: Dict[str, Dict[str, str]] = {}

    log("\nĐANG ĐỌC VÀ BÓC TÁCH PDF...")
    for idx, pdf in enumerate(pdf_files, 1):
        pdf_name = os.path.basename(pdf)
        try:
            text = read_pdf(pdf)
        except Exception as e:
            log(f"⚠️ Lỗi đọc file {pdf_name}: {e}")
            continue

        if "14-MST" in text or "14- MST" in text:
            data = parse_14(text, pdf_name)
            for cccd, value in data.items():
                if cccd not in all_results:
                    all_results[cccd] = value
                elif all_results[cccd]["status"] == "CAP_MOI":
                    continue
                elif value["status"] == "CAP_MOI":
                    all_results[cccd] = value
            log(f"[{idx}/{len(pdf_files)}] {pdf_name:<30} -> {len(data)} hồ sơ cấp MST")

        elif "40/TB" in text:
            data = parse_40(text, pdf_name)
            for cccd, value in data.items():
                if cccd not in all_results:
                    all_results[cccd] = value
                elif all_results[cccd]["status"] == "CAP_MOI":
                    continue
                elif value["status"] == "CAP_MOI":
                    all_results[cccd] = value
            log(f"[{idx}/{len(pdf_files)}] {pdf_name:<30} -> {len(data)} hồ sơ trùng MST")

        else:
            log(f"[{idx}/{len(pdf_files)}] {pdf_name:<30} -> Không nhận diện được mẫu")

        if progress_callback:
            progress_callback(idx, len(pdf_files))

    log(f"\n=> TỔNG SỐ HỒ SƠ LẤY ĐƯỢC TỪ PDF: {len(all_results)}")

    # Mở và cập nhật Excel
    log("\nĐang mở file Excel...")
    wb = load_workbook(excel_file)
    ws = wb.active

    last_col = ws.max_column
    status_col = get_column_letter(last_col + 1)
    note_col = get_column_letter(last_col + 2)
    source_col = get_column_letter(last_col + 3)

    log(f"Thêm cột kết quả: TRANG_THAI ({status_col}), GHI_CHU ({note_col}), FILE_NGUON ({source_col})")

    # Header ở dòng 3
    ws[f"{status_col}3"] = "TRANG_THAI"
    ws[f"{note_col}3"] = "GHI_CHU"
    ws[f"{source_col}3"] = "FILE_NGUON"

    green_fill = PatternFill(fill_type="solid", start_color="C6EFCE")
    yellow_fill = PatternFill(fill_type="solid", start_color="FFF2CC")

    updated = 0
    duplicate = 0
    total_data_rows = max(0, ws.max_row - 3)

    log(f"Đang đối chiếu dữ liệu Excel (tổng cộng {total_data_rows} dòng)...")

    for row in range(4, ws.max_row + 1):
        cccd_raw = ws[f"M{row}"].value
        if cccd_raw is None:
            continue

        cccd = str(cccd_raw)
        cccd = re.sub(r'\D', '', cccd).zfill(12)

        if cccd in all_results:
            result = all_results[cccd]

            # Cập nhật MST vào cột D
            ws[f"D{row}"] = result["mst"]
            ws[f"{status_col}{row}"] = result["status"]
            ws[f"{note_col}{row}"] = result["note"]
            ws[f"{source_col}{row}"] = result["source"]

            if result["status"] == "CAP_MOI":
                ws[f"D{row}"].fill = green_fill
            else:
                ws[f"D{row}"].fill = yellow_fill
                duplicate += 1

            updated += 1

    # Lưu file
    if not output_file:
        dir_name = os.path.dirname(excel_file)
        base_name = os.path.splitext(os.path.basename(excel_file))[0]
        output_file = os.path.join(dir_name, f"{base_name}_Updated.xlsx")

    wb.save(output_file)

    not_found = total_data_rows - updated
    log("\n==============================")
    log("HOÀN THÀNH XỬ LÝ MST")
    log("==============================")
    log(f"File kết quả   : {output_file}")
    log(f"Tổng số dòng   : {total_data_rows}")
    log(f"Tổng cập nhật  : {updated}")
    log(f"- Cấp mới (Xanh): {updated - duplicate}")
    log(f"- Trùng MST (Vàng): {duplicate}")
    log(f"Không tìm thấy : {not_found}")
    log("==============================\n")

    return {
        "output_file": output_file,
        "total_rows": total_data_rows,
        "updated": updated,
        "cap_moi": updated - duplicate,
        "duplicate": duplicate,
        "not_found": not_found
    }
