"""
Module xử lý Mã số thuế Người phụ thuộc (NPT)
Logic được chuyển đổi nguyên vẹn từ NPT.ipynb (Phiên bản 1.4)
"""

import os
import re
from typing import Callable, List, Optional, Dict, Any
try:
    import pymupdf as fitz
except ImportError:
    import fitz
from openpyxl import load_workbook
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter


def process_npt(
    excel_file: str,
    pdf_files: List[str],
    output_file: Optional[str] = None,
    log_callback: Optional[Callable[[str], None]] = None,
    progress_callback: Optional[Callable[[int, int], None]] = None
) -> Dict[str, Any]:
    """
    Thực thi quy trình đối soát & cập nhật file Người Phụ Thuộc (NPT).
    
    Args:
        excel_file: Đường dẫn file Excel nguồn (Tong hop NPT.xlsx)
        pdf_files: Danh sách đường dẫn file PDF thông báo từ Cục Thuế
        output_file: Đường dẫn file xuất kết quả (nếu None sẽ tự sinh)
        log_callback: Hàm nhận log message hiển thị lên UI
        progress_callback: Hàm nhận tiến độ (current, total)
        
    Returns:
        Dict chứa thống kê kết quả
    """
    def log(msg: str):
        if log_callback:
            log_callback(msg)
        else:
            print(msg)

    if not os.path.exists(excel_file):
        raise FileNotFoundError(f"Không tìm thấy file Excel: {excel_file}")
    if not pdf_files:
        raise ValueError("Chưa chọn file PDF nào để đối soát.")

    log(f"Bắt đầu xử lý Người Phụ Thuộc (NPT)...")
    log(f"- File Excel: {os.path.basename(excel_file)}")
    log(f"- Số lượng PDF: {len(pdf_files)} file")

    pdf_14_texts: Dict[str, str] = {}
    pdf_40_results: Dict[str, Dict[str, str]] = {}

    log("\nĐANG PHÂN TÍCH PDF...")
    for idx, pdf in enumerate(pdf_files, 1):
        pdf_name = os.path.basename(pdf)
        try:
            doc = fitz.open(pdf)
            text = ""
            for page in doc:
                # Đọc theo luồng gốc của PDF (không sort) để không làm nát bảng biểu
                text += page.get_text()
        except Exception as e:
            log(f"⚠️ Lỗi đọc file {pdf_name}: {e}")
            continue

        # Nếu là Mẫu 40 (Trùng MST - Không thành công)
        if "40/TB" in text or "Mẫu số: 40" in text:
            text_clean = re.sub(r'\s+', ' ', text)
            pattern = r'\b(\d{12})\b(?:(?!\b\d{12}\b).)*?MST(?:(?!\b\d{12}\b).)*?(\b\d{10}\b)'
            matches = re.findall(pattern, text_clean, flags=re.IGNORECASE | re.DOTALL)

            for cccd, mst in matches:
                pdf_40_results[cccd] = {
                    "mst": mst,
                    "status": "TRUNG_MST",
                    "note": f"MST đã tồn tại: {mst}",
                    "source": pdf_name
                }
            log(f"[{idx}/{len(pdf_files)}] {pdf_name:<30} -> {len(matches)} hồ sơ trùng MST")

        # Nếu là Mẫu 14 (Cấp mới - Thành công)
        elif "14-MST" in text or "14- MST" in text or "Mẫu số: 14" in text:
            # Thuật toán rút trích số: Xóa sạch mọi thứ không phải số
            text_only_digits = re.sub(r'\D', '', text)
            pdf_14_texts[pdf_name] = text_only_digits
            log(f"[{idx}/{len(pdf_files)}] {pdf_name:<30} -> Đã nạp dữ liệu cấp mới")

        else:
            log(f"[{idx}/{len(pdf_files)}] {pdf_name:<30} -> Không nhận diện được mẫu")

        if progress_callback:
            progress_callback(idx, len(pdf_files))

    # Cập nhật Excel
    log("\nĐANG MỞ VÀ CẬP NHẬT EXCEL...")
    wb = load_workbook(excel_file)
    ws = wb.active

    last_col = ws.max_column
    status_col = get_column_letter(last_col + 1)
    note_col = get_column_letter(last_col + 2)
    source_col = get_column_letter(last_col + 3)

    log(f"Thêm cột kết quả: TRANG_THAI ({status_col}), GHI_CHU ({note_col}), FILE_NGUON ({source_col})")

    # Header ở dòng 3
    ws[f"{status_col}3"] = "TRANG_THAI"
    ws[f"{note_col}3"] = "GHI_CHU"
    ws[f"{source_col}3"] = "FILE_NGUON"

    green_fill = PatternFill(fill_type="solid", start_color="C6EFCE")
    yellow_fill = PatternFill(fill_type="solid", start_color="FFF2CC")

    updated = 0
    duplicate = 0
    split_match_count = 0
    total_data_rows = max(0, ws.max_row - 3)

    log(f"Đang đối chiếu dữ liệu Excel (tổng cộng {total_data_rows} dòng, đọc CCCD tại Cột K)...")

    for row in range(4, ws.max_row + 1):
        cccd_val = ws[f"K{row}"].value
        if cccd_val is None:
            continue

        # Làm sạch CCCD gốc từ Excel
        cccd = re.sub(r'\D', '', str(cccd_val).replace(".0", "")).zfill(12)
        if not cccd:
            continue

        found = False

        # BƯỚC 1: TÌM TRONG MẪU 14 (ƯU TIÊN CẤP MỚI)
        for pdf_name, text_only_digits in pdf_14_texts.items():
            # 1. Dò bình thường
            if cccd in text_only_digits:
                ws[f"D{row}"] = cccd
                ws[f"{status_col}{row}"] = "CAP_MOI"
                ws[f"{note_col}{row}"] = "Được cấp MST"
                ws[f"{source_col}{row}"] = pdf_name
                ws[f"D{row}"].fill = green_fill
                updated += 1
                found = True
                break
            else:
                # 2. Khớp nối 2 mảnh phòng lỗi ngắt trang (Page Break)
                is_split_match = False
                for i in range(4, 9):  # Cắt thử ở vị trí 4 đến 8
                    part1 = cccd[:i]
                    part2 = cccd[i:]
                    pattern = f"{part1}\\d{{1,50}}{part2}"
                    if re.search(pattern, text_only_digits):
                        is_split_match = True
                        break

                if is_split_match:
                    ws[f"D{row}"] = cccd
                    ws[f"{status_col}{row}"] = "CAP_MOI"
                    ws[f"{note_col}{row}"] = "Được cấp MST (Khớp 2 mảnh do ngắt trang)"
                    ws[f"{source_col}{row}"] = pdf_name
                    ws[f"D{row}"].fill = green_fill
                    updated += 1
                    split_match_count += 1
                    found = True
                    break

        # BƯỚC 2: NẾU KHÔNG CÓ TRONG MẪU 14 -> TÌM MẪU 40
        if not found:
            if cccd in pdf_40_results:
                result = pdf_40_results[cccd]
                ws[f"D{row}"] = result["mst"]
                ws[f"{status_col}{row}"] = result["status"]
                ws[f"{note_col}{row}"] = result["note"]
                ws[f"{source_col}{row}"] = result["source"]
                ws[f"D{row}"].fill = yellow_fill
                duplicate += 1
                updated += 1
                found = True

    # Lưu file
    if not output_file:
        dir_name = os.path.dirname(excel_file)
        base_name = os.path.splitext(os.path.basename(excel_file))[0]
        output_file = os.path.join(dir_name, f"{base_name}_Updated.xlsx")

    wb.save(output_file)

    not_found = total_data_rows - updated
    log("\n==============================")
    log("HOÀN THÀNH XỬ LÝ NPT")
    log("==============================")
    log(f"File kết quả    : {output_file}")
    log(f"Tổng số dòng    : {total_data_rows}")
    log(f"Tổng cập nhật   : {updated}")
    log(f"- Cấp mới (12s) : {updated - duplicate} (trong đó {split_match_count} hồ sơ vá ngắt trang)")
    log(f"- Trùng MST     : {duplicate}")
    log(f"Không tìm thấy  : {not_found}")
    log("==============================\n")

    return {
        "output_file": output_file,
        "total_rows": total_data_rows,
        "updated": updated,
        "cap_moi": updated - duplicate,
        "split_matches": split_match_count,
        "duplicate": duplicate,
        "not_found": not_found
    }
